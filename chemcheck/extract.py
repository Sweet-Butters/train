"""PDF / PPTX에서 장(슬라이드)별 텍스트와 그림을 뽑는다.

같은 장에 있는 이름 텍스트와 구조 그림을 짝지어야 판정이 가능하므로
장 단위로 묶어서 돌려준다.

강의자료의 구조도는 두 경로로 들어온다.

  1. 래스터 이미지 — ChemDraw에서 복사해 붙인 그림
  2. 벡터 — PPT 도형으로 직접 그린 것

2번은 임베드된 이미지가 하나도 없다. 임베드 이미지만 보면 "그림 0개"로
조용히 넘어가는데, 그건 판정이 아니라 침묵이다. 그래서 최소한 못 뽑았다는
사실은 남긴다.
"""
from __future__ import annotations

from dataclasses import dataclass, field
from io import BytesIO
from pathlib import Path

# 아이콘·로고처럼 너무 작은 그림은 구조도가 아니다.
MIN_IMAGE_BYTES = 4096
MIN_IMAGE_SIDE = 80

# PPTX에서 구조로 인정할 최소 도형 수.
# 벤젠 고리를 선 도형으로 그리면 결합선만 6개가 된다.
MIN_PPTX_VECTOR_SHAPES = 6


@dataclass
class VectorRegion:
    """도형으로 그려졌지만 아직 이미지로 만들지 못한 구조 후보."""

    slide_index: int
    part_count: int
    grouped: bool


@dataclass
class Slide:
    index: int
    text: str
    images: list[Path] = field(default_factory=list)
    # 뽑지 못한 벡터 구조 후보. pipeline 은 읽지 않고 리포트가 쓴다.
    vector_regions: list[VectorRegion] = field(default_factory=list)


def _keep(width: int, height: int, size: int) -> bool:
    return size >= MIN_IMAGE_BYTES and min(width, height) >= MIN_IMAGE_SIDE


def _keep_blob(blob: bytes) -> bool:
    """바이트 크기와 실제 화소 크기를 모두 본다."""
    if len(blob) < MIN_IMAGE_BYTES:
        return False
    try:
        from PIL import Image

        with Image.open(BytesIO(blob)) as im:
            width, height = im.size
    except Exception:
        # 크기를 못 읽으면 버리지 않는다. 놓치는 쪽이 더 나쁘다.
        return True
    return min(width, height) >= MIN_IMAGE_SIDE


# ---------------------------------------------------------------------- PDF


def from_pdf(path: Path, out_dir: Path) -> list[Slide]:
    import pymupdf as fitz

    out_dir.mkdir(parents=True, exist_ok=True)
    slides: list[Slide] = []
    with fitz.open(path) as doc:
        for page_no, page in enumerate(doc, start=1):
            slide = Slide(page_no, page.get_text())
            for img_no, info in enumerate(page.get_images(full=True), start=1):
                xref = info[0]
                try:
                    pix = fitz.Pixmap(doc, xref)
                except Exception:
                    continue
                if pix.colorspace is None:
                    continue
                if pix.n - pix.alpha > 3:  # CMYK 등은 RGB로 변환
                    pix = fitz.Pixmap(fitz.csRGB, pix)
                data = pix.tobytes("png")
                if not _keep(pix.width, pix.height, len(data)):
                    continue
                dest = out_dir / f"p{page_no:03d}_{img_no:02d}.png"
                dest.write_bytes(data)
                slide.images.append(dest)
            slides.append(slide)
    return slides


# --------------------------------------------------------------------- PPTX


def _flatten(shapes, group_id: int | None = None):
    """그룹 안까지 평탄화한다. (도형, 최상위 그룹 id) 를 낸다.

    강의자료의 그림은 대개 설명 상자와 함께 그룹으로 묶여 있다.
    그룹을 안 들어가면 그 장은 통째로 그림 0개가 된다.
    """
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            top = group_id if group_id is not None else id(shape)
            yield from _flatten(shape.shapes, top)
        else:
            yield shape, group_id


def _shape_text(shape) -> str:
    """도형과 표에서 텍스트를 모은다."""
    parts: list[str] = []
    if getattr(shape, "has_text_frame", False):
        stripped = shape.text_frame.text.strip()
        if stripped:
            parts.append(stripped)
    if getattr(shape, "has_table", False):
        for row in shape.table.rows:
            for cell in row.cells:
                stripped = cell.text.strip()
                if stripped:
                    parts.append(stripped)
    return "\n".join(parts)


def _picture(shape):
    """그림이면 (blob, 확장자), 아니면 None.

    자리표시자 안에 들어간 그림은 shape_type 이 PICTURE 가 아니라서
    타입 대신 image 속성 유무로 판단한다.
    """
    try:
        image = shape.image
    except (AttributeError, ValueError, KeyError):
        return None
    return image.blob, (image.ext or "png")


def _is_vector_part(shape) -> bool:
    """구조를 이루는 도형 조각인가. 글자만 든 상자와 그림은 제외한다."""
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    kinds = {
        kind
        for kind in (
            getattr(MSO_SHAPE_TYPE, name, None)
            for name in ("AUTO_SHAPE", "FREEFORM", "LINE")
        )
        if kind is not None
    }
    if shape.shape_type not in kinds:
        return False
    return not _shape_text(shape)


def from_pptx(path: Path, out_dir: Path) -> list[Slide]:
    from pptx import Presentation

    out_dir.mkdir(parents=True, exist_ok=True)
    slides: list[Slide] = []
    deck = Presentation(str(path))
    for slide_no, raw in enumerate(deck.slides, start=1):
        texts: list[str] = []
        images: list[Path] = []
        img_no = 0
        # 그룹별 도형 조각 수. None 키는 그룹에 안 묶인 조각들.
        parts: dict[int | None, int] = {}

        for shape, group_id in _flatten(raw.shapes):
            stripped = _shape_text(shape)
            if stripped:
                texts.append(stripped)

            found = _picture(shape)
            if found is not None:
                blob, ext = found
                if not _keep_blob(blob):
                    continue
                img_no += 1
                dest = out_dir / f"s{slide_no:03d}_{img_no:02d}.{ext}"
                dest.write_bytes(blob)
                images.append(dest)
                continue

            if _is_vector_part(shape):
                parts[group_id] = parts.get(group_id, 0) + 1

        regions = [
            VectorRegion(slide_no, count, group_id is not None)
            for group_id, count in parts.items()
            if count >= MIN_PPTX_VECTOR_SHAPES
        ]
        slides.append(Slide(slide_no, "\n".join(texts), images, regions))
    return slides


def load(path: Path, out_dir: Path) -> list[Slide]:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return from_pdf(path, out_dir)
    if suffix == ".pptx":
        return from_pptx(path, out_dir)
    raise ValueError(f"지원하지 않는 형식: {suffix} (pdf 또는 pptx만 가능)")
