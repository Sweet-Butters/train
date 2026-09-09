"""심사위원이 눈으로 못 잡는 데모 슬라이드를 만든다.

의도적으로 틀린 구조를 넣는다. 현장에서 AI에게 즉석으로 시키면
맞게 나올 수도 있으므로 미리 고정해 둔다.

라벨은 여기서 정하지 않는다. bench/cases.py 의 CORPUS 가 정본이고 데모는 그
일부를 골라 쓴다. 라벨이 두 곳에 있으면 갈라지고, 갈라진 뒤에는 어느 쪽이
정본이었는지 아무도 모르게 된다.
"""
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from bench.cases import Truth, draw, pick  # noqa: E402

OUT = ROOT / "samples"

# 데모는 아세틸기 하나, 메틸기 하나 차이를 보여준다. 눈으로는 거의 구분되지 않는다.
DEMO = [
    pick("Aspirin", Truth.SAME),
    pick("Aspirin", Truth.SKELETON_DIFF),
    pick("Caffeine", Truth.SAME),
    pick("Caffeine", Truth.SKELETON_DIFF),
]

# 데모의 (맞다/틀리다) 두 갈래 형식은 '주의'(입체화학만 다름) 판정을 표현할 수 없다.
if any(c.truth is Truth.STEREO_DIFF for c in DEMO):
    raise ValueError("데모 형식으로는 '주의' 판정을 보여줄 수 없다. bench 로 볼 것")

# 기존 형태를 유지한다: (슬라이드 제목, 그려질 실제 구조 SMILES, 이게 맞는 구조인가)
CASES = [(c.name, c.smiles, c.truth is Truth.SAME) for c in DEMO]


def build() -> Path:
    from pptx import Presentation
    from pptx.util import Inches, Pt

    OUT.mkdir(parents=True, exist_ok=True)
    deck = Presentation()
    blank = deck.slide_layouts[6]

    for i, (name, smiles, _correct) in enumerate(CASES, start=1):
        png = OUT / f"demo_{i:02d}.png"
        draw(smiles, png)

        slide = deck.slides.add_slide(blank)
        box = slide.shapes.add_textbox(Inches(0.6), Inches(0.4), Inches(8.8), Inches(1.0))
        para = box.text_frame.paragraphs[0]
        para.text = name
        para.font.size = Pt(40)
        para.font.bold = True

        slide.shapes.add_picture(str(png), Inches(2.6), Inches(1.6), height=Inches(4.2))

    dest = OUT / "demo_slides.pptx"
    deck.save(str(dest))
    return dest


if __name__ == "__main__":
    path = build()
    print(f"생성: {path}")
    for i, (name, smiles, correct) in enumerate(CASES, start=1):
        mark = "정답" if correct else "의도적 오류"
        print(f"  {i}장  {name:10} {smiles:32} {mark}")
