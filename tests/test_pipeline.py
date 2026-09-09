"""판정 파이프라인 검증.

실제 OCSR 없이도 판정 로직을 검증하기 위해 '무엇이 그려졌는지 아는' 스텁
인식기를 쓴다. 이 스텁은 완벽한 인식기를 흉내낸 것이며 제품 경로가 아니다.
"""
from __future__ import annotations

import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from chemcheck.ocsr import Engine, Prediction  # noqa: E402
from chemcheck.pipeline import run  # noqa: E402
from chemcheck.verdict import Verdict  # noqa: E402
from scripts.make_demo import CASES, build  # noqa: E402


class OracleEngine(Engine):
    """그려진 SMILES를 그대로 돌려주는 스텁 = 정확도 100%인 가상 인식기."""

    name = "oracle"

    def __init__(self, by_slide: dict[int, str]):
        self.by_slide = by_slide

    def available(self) -> bool:
        return True

    def recognize(self, image_path: Path) -> Prediction | None:
        slide_no = int(image_path.stem.split("_")[0].lstrip("s"))
        smiles = self.by_slide.get(slide_no)
        return Prediction(smiles, 0.99, self.name) if smiles else None


def test_detects_planted_errors() -> None:
    oracle = OracleEngine({i: smiles for i, (_, smiles, _) in enumerate(CASES, start=1)})
    expected = [Verdict.OK if correct else Verdict.ERROR for _, _, correct in CASES]

    with tempfile.TemporaryDirectory() as tmp:
        deck = build(Path(tmp))
        results = run(deck, Path(tmp) / "work", [oracle])

    got = [finding.verdict for r in results for _, finding in r.findings]
    assert got == expected, f"기대 {expected}, 실제 {got}"

    for r, want in zip(results, expected):
        for _, finding in r.findings:
            print(f"  {r.slide.index}장 {r.references[0].name:9} -> "
                  f"{finding.verdict.value.upper():6} {finding.reason}")
    print("\n통과: 심어둔 오류 2건을 모두 잡고 정답 2건은 통과시켰다.")




class FlakyEngine(Engine):
    """이미지를 조금만 바꿔도 답이 흔들리는 인식기. 실제 오인식 상황을 흉내낸다."""

    name = "flaky"

    def __init__(self):
        self.calls = 0

    def available(self) -> bool:
        return True

    def recognize(self, image_path: Path) -> Prediction | None:
        self.calls += 1
        # 첫 호출은 아스피린, 이후 변형본에서는 카페인 — 즉 스스로 합의하지 못한다.
        smiles = "CC(=O)Oc1ccccc1C(=O)O" if self.calls % 3 == 1 else "CN1C=NC2=C1C(=O)N(C(=O)N2C)C"
        return Prediction(smiles, float("nan"), self.name)


def test_self_consistency_forces_abstain() -> None:
    """인식이 흔들리면 '오류'라고 말하지 않고 보류해야 한다."""
    from chemcheck.ocsr import SelfConsistent

    engine = SelfConsistent(FlakyEngine())

    with tempfile.TemporaryDirectory() as tmp:
        deck = build(Path(tmp))
        results = run(deck, Path(tmp) / "work", [engine])

    verdicts = [f.verdict for r in results for _, f in r.findings]
    assert all(v is Verdict.ABSTAIN for v in verdicts), f"보류가 아님: {verdicts}"

    reason = results[0].findings[0][1].reason
    print(f"  흔들리는 인식 -> {verdicts[0].value.upper()}: {reason}")
    print("통과: 인식이 스스로 합의하지 못하면 오류로 단정하지 않는다.")


class StemOracle(Engine):
    """그림 파일 이름별로 '실제로 그려진 것'을 아는 완벽한 인식기."""

    name = "stem-oracle"

    def __init__(self, by_stem: dict[str, str]):
        self.by_stem = by_stem

    def available(self) -> bool:
        return True

    def recognize(self, image_path: Path) -> Prediction | None:
        smiles = self.by_stem.get(image_path.stem)
        return Prediction(smiles, 0.99, self.name) if smiles else None


SMILES = {
    "aspirin": "CC(=O)Oc1ccccc1C(=O)O",
    "salicylic acid": "OC(=O)c1ccccc1O",
    "caffeine": "CN1C=NC2=C1C(=O)N(C(=O)N2C)C",
    "theobromine": "CN1C=NC2=C1C(=O)NC(=O)N2C",
    "paracetamol": "CC(=O)Nc1ccc(O)cc1",
    "ibuprofen": "CC(C)Cc1ccc(cc1)C(C)C(=O)O",
}


def _grid_deck(dest: Path, title: str, cells: list[tuple[str, str]]) -> None:
    """한 장에 (라벨, 실제로 그려진 분자)를 격자로 배치한 덱을 만든다."""
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from rdkit import Chem
    from rdkit.Chem import Draw

    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    head = slide.shapes.add_textbox(Inches(0.4), Inches(0.2), Inches(9.0), Inches(0.8))
    head.text_frame.paragraphs[0].text = title
    head.text_frame.paragraphs[0].runs[0].font.size = Pt(24)

    for i, (label, actual) in enumerate(cells):
        png = dest.parent / f"cell_{i}.png"
        Draw.MolToFile(Chem.MolFromSmiles(SMILES[actual]), str(png), size=(300, 240))
        left = Inches(0.4 + 4.6 * (i % 2))
        top = Inches(1.2 + 2.7 * (i // 2))
        slide.shapes.add_picture(str(png), left, top, height=Inches(2.2))
        cap = slide.shapes.add_textbox(left, top + Inches(2.25), Inches(4.2), Inches(0.4))
        cap.text_frame.paragraphs[0].text = label
    deck.save(str(dest))


def _run_grid(tmp: Path, title: str, cells: list[tuple[str, str]]) -> list[Verdict]:
    from chemcheck.extract import load

    deck = tmp / "grid.pptx"
    _grid_deck(deck, title, cells)
    work = tmp / "work"
    stems = [p.stem for p in load(deck, work)[0].images]
    oracle = StemOracle({stem: SMILES[actual] for stem, (_, actual) in zip(stems, cells)})
    results = run(deck, work, [oracle])
    return [f.verdict for r in results for _, f in r.findings]


def test_multi_structure_slide_is_paired_by_position() -> None:
    """한 장에 구조가 여럿일 때, 그림은 제 캡션하고만 대조해야 한다.

    장에 있는 아무 이름에나 맞춰보면 뒤바뀐 그림이 통과해 버린다.
    """
    with tempfile.TemporaryDirectory() as tmp:
        got = _run_grid(
            Path(tmp),
            "Stimulants: Caffeine vs Theobromine",
            [("Caffeine", "theobromine"), ("Theobromine", "caffeine")],
        )
    assert all(v is Verdict.ERROR for v in got), f"뒤바뀐 그림을 놓쳤다: {got}"
    print(f"  뒤바뀐 구조 2건 -> {[v.value.upper() for v in got]}")
    print("통과: 옆 그림의 이름과 우연히 맞아도 통과시키지 않는다.")


def test_multi_structure_slide_has_no_false_error() -> None:
    """전부 올바른 장에서는 하나도 오류라고 말하지 않아야 한다."""
    with tempfile.TemporaryDirectory() as tmp:
        got = _run_grid(
            Path(tmp),
            "Analgesics overview",
            [("Aspirin", "aspirin"), ("Paracetamol", "paracetamol"),
             ("Ibuprofen", "ibuprofen"), ("Salicylic acid", "salicylic acid")],
        )
    assert Verdict.ERROR not in got, f"멀쩡한 그림을 틀렸다고 했다: {got}"
    assert all(v is Verdict.OK for v in got), f"과도한 보류: {got}"
    print(f"  정상 구조 4건 -> {[v.value.upper() for v in got]}")
    print("통과: 이름이 여럿이어도 거짓 오류를 내지 않는다.")


if __name__ == "__main__":
    test_detects_planted_errors()
    print()
    test_self_consistency_forces_abstain()
    print()
    test_multi_structure_slide_is_paired_by_position()
    print()
    test_multi_structure_slide_has_no_false_error()
