"""추출 검증.

인식기 없이 돌아간다. 강의자료가 구조를 담는 두 방식(래스터/벡터)을
실제 파일로 만들어서, 뽑히는 것과 뽑히지 말아야 할 것을 나눈다.

가장 중요한 두 가지:
  - 벡터로 그린 구조가 "그림 0개"로 조용히 사라지지 않는다
  - 그룹 안에 든 그림이 누락되지 않는다
"""
from __future__ import annotations

import math
import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from chemcheck import extract  # noqa: E402
from chemcheck.extract import from_pdf, from_pptx, load  # noqa: E402

PAGE_W, PAGE_H = 612, 792


# ------------------------------------------------------------------ 픽스처


def _pdf(dest: Path, pages: list[bytes]) -> Path:
    """내용 스트림만 받아 최소 PDF를 만든다. 외부 라이브러리를 쓰지 않는다."""
    font_id = 3 + 2 * len(pages)
    kids = " ".join(f"{3 + 2 * i} 0 R" for i in range(len(pages)))
    objs: list[bytes] = [
        b"<</Type/Catalog/Pages 2 0 R>>",
        f"<</Type/Pages/Kids[{kids}]/Count {len(pages)}>>".encode(),
    ]
    for i, content in enumerate(pages):
        objs.append(
            f"<</Type/Page/Parent 2 0 R/MediaBox[0 0 {PAGE_W} {PAGE_H}]"
            f"/Contents {4 + 2 * i} 0 R"
            f"/Resources<</Font<</F1 {font_id} 0 R>>>>>>".encode()
        )
        objs.append(b"<</Length %d>>\nstream\n" % len(content) + content + b"\nendstream")
    objs.append(b"<</Type/Font/Subtype/Type1/BaseFont/Helvetica>>")

    out = bytearray(b"%PDF-1.4\n")
    offsets: list[int] = []
    for i, body in enumerate(objs, start=1):
        offsets.append(len(out))
        out += b"%d 0 obj " % i + body + b" endobj\n"
    xref_at = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objs) + 1)
    for off in offsets:
        out += b"%010d 00000 n \n" % off
    out += b"trailer <</Size %d/Root 1 0 R>>\nstartxref\n%d\n%%%%EOF\n" % (
        len(objs) + 1,
        xref_at,
    )
    dest.write_bytes(bytes(out))
    return dest


def _ring(x: float, y: float, r: float = 40) -> bytes:
    """육각 고리 하나를 경로 한 개로 그린다. ChemDraw 내보내기와 같은 모양."""
    pts = [
        (x + r * math.cos(math.radians(60 * i)), y + r * math.sin(math.radians(60 * i)))
        for i in range(6)
    ]
    head = f"{pts[0][0]:.1f} {pts[0][1]:.1f} m "
    rest = " ".join(f"{px:.1f} {py:.1f} l" for px, py in pts[1:])
    return (head + rest + " h S\n").encode()


def _title(text: str, y: float = 720) -> bytes:
    return f"BT /F1 24 Tf 72 {y} Td ({text}) Tj ET\n".encode()


def _corner(x: float, y: float, size: float = 100) -> bytes:
    """ㄱ자 두 획. 크기는 충분하지만 꼭짓점이 모자라 구조가 아니다."""
    return (
        f"{x} {y} m {x + size} {y} l S\n{x} {y} m {x} {y + size} l S\n"
    ).encode()


def _png(dest: Path, side: int) -> Path:
    """실제 png 파일. 노이즈를 넣어 압축돼도 최소 크기를 넘기게 한다."""
    import random

    from PIL import Image

    rng = random.Random(0)
    image = Image.new("RGB", (side, side))
    image.putdata([
        (rng.randrange(256), rng.randrange(256), rng.randrange(256))
        for _ in range(side * side)
    ])
    image.save(dest)
    return dest


def _ink(path: Path) -> int:
    from PIL import Image

    with Image.open(path) as im:
        return sum(1 for v in im.convert("L").tobytes() if v < 200)


# ---------------------------------------------------------------- PDF: 벡터


def test_pdf_vector_structures_are_extracted() -> None:
    """도형으로 그린 구조 2개가 각각 그림으로 나온다.

    임베드 이미지가 0개인 페이지다. 예전 방식이라면 전부 놓쳤다.
    """
    with tempfile.TemporaryDirectory() as tmp:
        work = Path(tmp)
        deck = _pdf(work / "vector.pdf", [_title("Aspirin") + _ring(200, 500) + _ring(450, 300)])

        slides = from_pdf(deck, work / "out")

        assert len(slides) == 1
        assert len(slides[0].images) == 2, f"구조 2개를 기대했는데 {len(slides[0].images)}개"
        for image in slides[0].images:
            assert image.exists()
            assert _ink(image) > 0, f"{image.name} 이 비어 있다"


def test_pdf_text_is_extracted() -> None:
    with tempfile.TemporaryDirectory() as tmp:
        work = Path(tmp)
        deck = _pdf(work / "vector.pdf", [_title("Aspirin") + _ring(200, 500)])

        slides = from_pdf(deck, work / "out")

        assert "Aspirin" in slides[0].text


def test_pdf_sparse_strokes_are_ignored() -> None:
    """크기는 충분해도 획이 두 개뿐이면 구조가 아니다."""
    with tempfile.TemporaryDirectory() as tmp:
        work = Path(tmp)
        deck = _pdf(work / "sparse.pdf", [_title("Table") + _corner(200, 400)])

        slides = from_pdf(deck, work / "out")

        assert slides[0].images == []


def test_pdf_rectangle_outline_is_ignored() -> None:
    """표 테두리·글상자 윤곽은 커도 구조가 아니다."""
    with tempfile.TemporaryDirectory() as tmp:
        work = Path(tmp)
        deck = _pdf(work / "rect.pdf", [_title("Table") + b"100 300 300 200 re S\n"])

        slides = from_pdf(deck, work / "out")

        assert slides[0].images == []


def test_pdf_pages_are_numbered_from_one() -> None:
    with tempfile.TemporaryDirectory() as tmp:
        work = Path(tmp)
        deck = _pdf(
            work / "two.pdf",
            [_title("Aspirin") + _ring(200, 500), _title("Caffeine") + _ring(200, 500)],
        )

        slides = from_pdf(deck, work / "out")

        assert [s.index for s in slides] == [1, 2]
        assert all(s.images for s in slides)


# --------------------------------------------------------------- PPTX: 그룹


def _deck(dest: Path, build) -> Path:
    from pptx import Presentation

    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    build(slide)
    deck.save(str(dest))
    return dest


def test_pptx_picture_inside_group_is_found() -> None:
    """그룹 안에 든 그림을 놓치지 않는다.

    강의자료의 구조도는 설명 상자와 함께 묶여 있는 경우가 많다.
    그룹을 안 들어가면 그 장은 통째로 그림 0개가 된다.
    """
    from pptx.util import Inches

    with tempfile.TemporaryDirectory() as tmp:
        work = Path(tmp)
        png = _png(work / "structure.png", 240)

        def build(slide):
            group = slide.shapes.add_group_shape()
            group.shapes.add_picture(str(png), Inches(1), Inches(1), Inches(3), Inches(3))

        deck = _deck(work / "grouped.pptx", build)
        slides = from_pptx(deck, work / "out")

        assert len(slides[0].images) == 1, "그룹 안의 그림을 놓쳤다"


def test_pptx_text_inside_group_is_found() -> None:
    from pptx.util import Inches

    with tempfile.TemporaryDirectory() as tmp:
        work = Path(tmp)

        def build(slide):
            group = slide.shapes.add_group_shape()
            box = group.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
            box.text_frame.text = "Aspirin"

        deck = _deck(work / "grouptext.pptx", build)
        slides = from_pptx(deck, work / "out")

        assert "Aspirin" in slides[0].text


def test_pptx_table_text_is_found() -> None:
    from pptx.util import Inches

    with tempfile.TemporaryDirectory() as tmp:
        work = Path(tmp)

        def build(slide):
            table = slide.shapes.add_table(
                2, 2, Inches(1), Inches(1), Inches(6), Inches(2)
            ).table
            table.cell(0, 0).text = "Caffeine"

        deck = _deck(work / "table.pptx", build)
        slides = from_pptx(deck, work / "out")

        assert "Caffeine" in slides[0].text


def test_pptx_small_image_is_filtered() -> None:
    """아이콘 크기 그림은 구조도가 아니다. 예전에는 바이트 수만 봤다."""
    from pptx.util import Inches

    with tempfile.TemporaryDirectory() as tmp:
        work = Path(tmp)
        icon = _png(work / "icon.png", 32)

        def build(slide):
            slide.shapes.add_picture(str(icon), Inches(1), Inches(1), Inches(1), Inches(1))

        deck = _deck(work / "icon.pptx", build)
        slides = from_pptx(deck, work / "out")

        assert slides[0].images == []


def test_pptx_drawn_structure_is_recorded() -> None:
    """PPT 도형으로 그린 구조는 아직 이미지로 못 만든다.

    그래도 못 뽑았다는 사실은 남긴다. 조용히 0개로 넘어가면
    사용자는 검사가 된 줄 안다.
    """
    from pptx.enum.shapes import MSO_CONNECTOR
    from pptx.util import Inches

    with tempfile.TemporaryDirectory() as tmp:
        work = Path(tmp)

        def build(slide):
            group = slide.shapes.add_group_shape()
            for i in range(6):
                x = Inches(2 + i * 0.3)
                group.shapes.add_connector(
                    MSO_CONNECTOR.STRAIGHT, x, Inches(2), x + Inches(0.3), Inches(2.5)
                )

        deck = _deck(work / "drawn.pptx", build)
        slides = from_pptx(deck, work / "out")

        assert slides[0].images == []
        assert len(slides[0].vector_regions) == 1
        region = slides[0].vector_regions[0]
        assert region.part_count >= extract.MIN_PPTX_VECTOR_SHAPES
        assert region.grouped is True


def test_pptx_few_shapes_are_not_a_structure() -> None:
    """화살표 한두 개를 구조라고 부르면 안 된다."""
    from pptx.enum.shapes import MSO_CONNECTOR
    from pptx.util import Inches

    with tempfile.TemporaryDirectory() as tmp:
        work = Path(tmp)

        def build(slide):
            slide.shapes.add_connector(
                MSO_CONNECTOR.STRAIGHT, Inches(1), Inches(1), Inches(3), Inches(1)
            )

        deck = _deck(work / "arrow.pptx", build)
        slides = from_pptx(deck, work / "out")

        assert slides[0].vector_regions == []


# ------------------------------------------------------------------- 공통


def test_load_rejects_unknown_suffix() -> None:
    with tempfile.TemporaryDirectory() as tmp:
        work = Path(tmp)
        stray = work / "notes.txt"
        stray.write_text("x", encoding="utf-8")

        try:
            load(stray, work / "out")
        except ValueError as err:
            assert "지원하지 않는 형식" in str(err)
        else:
            raise AssertionError("ValueError 를 기대했다")
