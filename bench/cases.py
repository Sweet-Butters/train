"""벤치마크 코퍼스 - 라벨 붙은 (이름, 실제로 그려진 구조) 쌍.

라벨은 사람이 선언하지만 믿지 않는다. `validate` 가 PubChem + RDKit 으로
같은 라벨을 다시 계산해 대조한다. 선언과 계산이 어긋나면 그건 도구가 아니라
케이스가 틀린 것이다. 이 대조는 OCSR 없이 오늘 돌아간다.
"""
from __future__ import annotations

from dataclasses import dataclass
from enum import Enum
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
DECKS = Path(__file__).resolve().parent / "decks"


class Truth(Enum):
    """이름과 그림의 실제 관계. 도구의 판정이 아니라 사실이다."""

    SAME = "same"                    # 같은 분자 - 여기서 ERROR 가 나오면 오탐
    SKELETON_DIFF = "skeleton_diff"  # 골격이 다름 - 잡아야 할 진짜 오류
    STEREO_DIFF = "stereo_diff"      # 골격은 같고 입체화학만 다름
    NOT_A_STRUCTURE = "not_structure"  # 애초에 구조식이 아님 - 판정하면 안 된다
    NO_CLAIM = "no_claim"              # 구조식은 맞지만 슬라이드가 이름을 대지 않았다


# 판정했어야 할 케이스. 아래 둘은 여기 들지 않는다.
JUDGEABLE = (Truth.SAME, Truth.SKELETON_DIFF, Truth.STEREO_DIFF)

# 판정하면 안 되는 케이스. 물러나는 것이 정답이다.
#   NOT_A_STRUCTURE - 그림 쪽에 검사 대상이 없다 (클립아트·사진·오비탈 도해)
#   NO_CLAIM        - 이름 쪽에 검사할 주장이 없다 (슬라이드가 화합물을 지목하지 않음)
# 둘 다 '대조할 짝이 없다'는 같은 사실의 양쪽이므로 채점에서 같이 다룬다.
NOT_JUDGEABLE = (Truth.NOT_A_STRUCTURE, Truth.NO_CLAIM)


@dataclass(frozen=True)
class Case:
    name: str      # 슬라이드에 적히는 이름
    smiles: str    # 슬라이드에 실제로 그려지는 구조 (구조식이 아니면 빈 문자열)
    truth: Truth
    note: str = ""
    art: str = ""    # 구조식이 아닌 그림의 종류. NOT_A_STRUCTURE 에서만 쓴다
    extra: str = ""  # 이름 옆에 함께 적히는 다른 글자. 참조를 흐리는 말들


# 합성 코퍼스: RDKit 이 그린 그림이다.
#
# 한계를 분명히 해 둔다 - 이 그림들은 전부 같은 렌더러에서 나온다. 실제 슬라이드는
# ChemDraw 내보내기, 스크린샷, 손그림이 섞여 있고 해상도도 제각각이다. 여기서 나온
# 숫자는 OCSR 의 상한이지 현장 성능이 아니다. 현장 숫자는 decks/ 의 실제 덱에서 나온다.
CORPUS: list[Case] = [
    # --- 같은 분자 (오탐 측정용) -------------------------------------------
    Case("Aspirin", "CC(=O)Oc1ccccc1C(=O)O", Truth.SAME),
    Case("Caffeine", "CN1C=NC2=C1C(=O)N(C(=O)N2C)C", Truth.SAME),
    Case("Paracetamol", "CC(=O)Nc1ccc(O)cc1", Truth.SAME),
    Case("Benzoic acid", "OC(=O)c1ccccc1", Truth.SAME),
    Case("Glucose", "OC[C@H]1OC(O)[C@H](O)[C@@H](O)[C@@H]1O", Truth.SAME,
         "고리·입체가 많아 OCSR 이 흔들리기 쉬운 쪽"),

    # --- 골격이 다름 (검출 측정용) -----------------------------------------
    Case("Aspirin", "OC(=O)c1ccccc1O", Truth.SKELETON_DIFF,
         "살리실산 - 아세틸기 하나 차이. 눈으로는 거의 구분되지 않는다"),
    Case("Caffeine", "CN1C=NC2=C1C(=O)NC(=O)N2C", Truth.SKELETON_DIFF,
         "테오브로민 - 메틸기 하나가 없다"),
    Case("Paracetamol", "CCOc1ccc(NC(C)=O)cc1", Truth.SKELETON_DIFF,
         "페나세틴 - 수산기가 에톡시로 바뀌었다"),

    # --- 입체화학만 다름 (등급 측정용) --------------------------------------
    Case("L-Alanine", "N[C@H](C)C(=O)O", Truth.STEREO_DIFF,
         "D-알라닌을 그렸다. 골격은 같으므로 오류가 아니라 주의여야 한다"),

    # --- 구조식이 아닌 그림 (물러남 측정용) ---------------------------------
    # 실제 강의자료에서 뽑은 그림 371개 중 고유한 것이 109개였고, 그중 구조식은
    # 소수였다. 나머지는 클립아트·오비탈 도해·실험기구·사진이다. 이름이 적힌 장에
    # 그런 그림이 섞여 있으면 도구는 그것도 인식기에 먹인다. 거기서 나온 SMILES 가
    # 이름과 다르면 '오류'가 되는데, 자료는 멀쩡하다. 실제 자료의 가장 큰 오탐원이
    # 여기일 수 있다. 아래는 그 경로를 재기 위한 자리표시 그림이다 - 진짜 클립아트는
    # 저작물이라 담지 않는다. 현장 숫자는 실제 덱에서 나온다.
    Case("Methane", "", Truth.NOT_A_STRUCTURE, "장식용 도형", art="shapes"),
    Case("Ethane", "", Truth.NOT_A_STRUCTURE, "글자만 있는 그림", art="text"),
    Case("Benzene", "", Truth.NOT_A_STRUCTURE, "사진 비슷한 잡음", art="noise"),

    # --- 참조를 흐리는 글자가 함께 있는 장 (오탐 측정용) ---------------------
    # 실제 자료에서 판정 후보 장 128개 중 24개(그림 77건)는 잡힌 참조가 전부
    # 분류명이나 잡음이었다. 'Alkenes' 가 47번, 'Sol' 14번, 'Challenge' 9번,
    # 'ene' 6번. PubChem 은 분류명과 상표명을 정식으로 들고 있어서 해석해 준다.
    # 그림은 맞게 그려졌는데 참조가 엉뚱하면 '오류'가 난다 - 오탐이다.
    Case("But-2-ene", "CC=CC", Truth.SAME,
         "장에 분류명 Alkenes 가 함께 있다. 참조가 그쪽으로 잡히면 오탐이 난다",
         extra="Alkenes"),
    Case("Propane", "CCC", Truth.SAME,
         "장에 Challenge 와 Sol 이 함께 있다. 둘 다 PubChem 이 화합물로 준다",
         extra="Challenge Sol"),
]


def pick(name: str, truth: Truth) -> Case:
    """코퍼스에서 케이스 하나를 집는다. 번호로 집으면 코퍼스가 늘 때 조용히 어긋난다."""
    for case in CORPUS:
        if case.name == name and case.truth is truth:
            return case
    raise KeyError(f"코퍼스에 없는 케이스: {name} / {truth.value}")


def draw(smiles: str, dest: Path) -> None:
    """SMILES 를 그림으로. scripts/make_demo.py 와 같은 렌더 경로를 쓴다."""
    from rdkit import Chem
    from rdkit.Chem import Draw

    mol = Chem.MolFromSmiles(smiles)
    if mol is None:
        raise ValueError(f"SMILES 파싱 실패: {smiles}")
    Draw.MolToFile(mol, str(dest), size=(500, 400))


def draw_non_structure(kind: str, dest: Path) -> None:
    """구조식이 아닌 그림을 만든다. 실제 클립아트의 자리표시다.

    진짜 클립아트·사진은 저작물이라 저장소에 담지 않는다. 여기서 재려는 것은
    '구조식이 아닌 그림을 인식기에 먹였을 때 도구가 물러나는가'라는 경로이지
    특정 그림이 아니다.
    """
    import random

    from PIL import Image, ImageDraw

    img = Image.new("RGB", (500, 400), "white")
    d = ImageDraw.Draw(img)
    rng = random.Random(kind)  # 같은 종류는 늘 같은 그림

    # extract.py 는 4096 바이트 미만 그림을 버린다. 단색 배경 위의 도형·글자는
    # 너무 잘 압축돼 그 밑으로 떨어져 추출 단계에서 조용히 사라진다. 옅은 얼룩을
    # 깔아 실제 그림의 크기를 흉내낸다 - 재려는 것은 압축률이 아니다.
    px = img.load()
    for y in range(0, 400, 3):
        for x in range(0, 500, 3):
            v = 235 + rng.randrange(20)
            px[x, y] = (v, v, v)

    if kind == "shapes":
        for _ in range(6):
            x, y = rng.randrange(40, 380), rng.randrange(40, 300)
            d.ellipse([x, y, x + 90, y + 70], fill=(rng.randrange(256), 90, 160))
    elif kind == "text":
        for i in range(6):
            d.text((40, 40 + i * 45), "Chapter summary and key points", fill="black")
    else:  # noise - 사진 비슷한 것
        for y in range(0, 400, 2):
            for x in range(0, 500, 2):
                v = (rng.randrange(256), rng.randrange(256), rng.randrange(256))
                for dy in (0, 1):
                    for dx in (0, 1):
                        px[x + dx, y + dy] = v
    img.save(dest)


def build_deck(cases: list[Case], dest: Path) -> Path:
    """케이스 하나당 한 장. 장 번호가 곧 케이스 번호가 되도록 한 장에 하나만 둔다."""
    from pptx import Presentation
    from pptx.util import Inches, Pt

    dest.parent.mkdir(parents=True, exist_ok=True)
    images = dest.parent / f"{dest.stem}_images"
    images.mkdir(parents=True, exist_ok=True)

    deck = Presentation()
    blank = deck.slide_layouts[6]
    for i, case in enumerate(cases, start=1):
        png = images / f"{i:02d}.png"
        if case.truth is Truth.NOT_A_STRUCTURE:
            draw_non_structure(case.art or "shapes", png)
        else:
            draw(case.smiles, png)

        slide = deck.slides.add_slide(blank)
        box = slide.shapes.add_textbox(Inches(0.6), Inches(0.4), Inches(8.8), Inches(1.0))
        para = box.text_frame.paragraphs[0]
        para.text = f"{case.name} {case.extra}".strip()
        para.font.size = Pt(40)
        para.font.bold = True
        slide.shapes.add_picture(str(png), Inches(2.6), Inches(1.6), height=Inches(4.2))

    deck.save(str(dest))
    return dest
